"""Formula export and actual engine recalculate/save/reopen checks."""
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
try:
    from .runtime import node_executable,bundled
except ImportError:
    from runtime import node_executable,bundled


def write_json(path,data):
    Path(path).write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False)+'\n')


def specification(data, graph):
    tabs=[]
    for name,paths in graph.sheets.items():
        rows=[dict(row=1,values=[data['company']+' - '+name]),
              dict(row=2,values=['SYNTHETIC / UNREVIEWED' if data['synthetic'] else 'CONDITIONAL / UNREVIEWED']),
              dict(row=3,values=[f"{data['currency']} amounts / scale {data['amount_scale']['value']:g}; valuation date {data['valuation_date']}"]),
              dict(row=4,values=['Calculation / input','Value','Unit','Result path / source','As of','Basis','Review','Rationale'])]
        for path in paths:
            node=graph.nodes[path]; meta=node['metadata']
            source=meta['source']['id']+' | '+meta['source']['locator'] if meta else path
            row=dict(row=int(node['cell'][1:]),unit=node['unit'],expected=node['value'],
                     values=[node['label'],node['value'],node['unit'],source,
                             meta['as_of'] if meta else '',meta['kind'] if meta else 'calculated',
                             json.dumps(meta['review']) if meta else 'not independently reviewed',
                             meta['rationale'] if meta else ''])
            if not meta: row['formula']='='+graph.formula(node['expression'])
            rows.append(row)
        tabs.append(dict(name=name,rows=rows))
    # Policy metadata is retained alongside numeric provenance, not silently dropped.
    from importlib import import_module
    leaf_fn=import_module((__package__+'.' if __package__ else '')+'schema_tools').leaves
    policies=[dict(row=1,values=['Metadata reference (not model drivers)']),
              dict(row=2,values=['Unreviewed input declarations; not verified legal or accounting treatment']),
              dict(row=4,values=['Input','Value','Unit','Source URI','Source SHA256','Review','Locator','Rationale'])]
    for row,(path,item) in enumerate(leaf_fn(data),5):
        policies.append(dict(row=row,values=[path,item['value'],item['unit'],item['source']['uri'],item['source']['sha256'],
                                             json.dumps(item['review']),item['source']['locator'],item['rationale']]))
    tabs.append(dict(name='Evidence',rows=policies))
    return dict(sheets=tabs)


def openpyxl_export(spec,out):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.workbook.properties import CalcProperties
    wb=Workbook();wb.remove(wb.active)
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    for tab in spec['sheets']:
        ws=wb.create_sheet(tab['name']);ws.sheet_view.showGridLines=False;ws.freeze_panes='B5'
        widths={'A':55,'B':24,'C':24,'D':50,'E':30,'F':35,'G':45,'H':70}
        for col,width in widths.items():ws.column_dimensions[col].width=width
        for row in tab['rows']:
            r=row['row'];ws.row_dimensions[r].height=23
            for c,value in enumerate(row['values'],1):
                cell=ws.cell(r,c,value)
                # Never allow evidence text / identifiers to become injected formulas.
                if isinstance(value,str):cell.data_type='s'
                cell.font=Font(name='Arial',size=10)
            if row.get('formula'):
                ws.cell(r,2,row['formula']).font=Font(name='Arial',size=10,color='008000')
            elif r>=5 and isinstance(row['values'][1],(int,float)):
                ws.cell(r,2).font=Font(name='Arial',size=10,color='0000FF')
            if r>=5:
                ws.cell(r,2).number_format='0.00%' if row.get('unit')=='ratio' else '#,##0.00;(#,##0.00);"-"'
            if tab['name']=='Evidence' and r>=5:
                ws.cell(r,2).font=Font(name='Arial',size=10,color='444444')
                if isinstance(row['values'][1],str):
                    ws.cell(r,2).alignment=Alignment(wrap_text=True,vertical='center')
                    ws.row_dimensions[r].height=max(23,math.ceil(len(row['values'][1])/26)*14+8)
            if tab['name']=='Checks' and r>=5:ws.cell(r,2).number_format='#,##0.00;(#,##0.00);0.00'
            if r==4:
                for cell in ws[r]:
                    cell.fill=PatternFill('solid',fgColor='193C47');cell.font=Font(name='Arial',color='FFFFFF',bold=True)
        ws.auto_filter.ref=f'A4:H{ws.max_row}'
    wb.save(out/'valuation.xlsx')


def artifact_location():
    override=os.environ.get('VALUATION_ARTIFACT_TOOL')
    if override:return Path(override)
    # Host runtime discovery, never an author checkout dependency.
    root=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules/@oai/artifact-tool'
    if (root/'package.json').exists():
        package=json.loads((root/'package.json').read_text())
        main=package.get('main','dist/artifact_tool.mjs')
        if (root/main).exists():return root/main
    return None


def export(data,graph,out,engine='auto'):
    spec=specification(data,graph)
    write_json(out/'workbook-spec.json',spec)
    artifact=artifact_location()
    node=node_executable()
    if engine=='artifact' or (engine=='auto' and artifact and node):
        env=os.environ.copy()
        if artifact:env['VALUATION_ARTIFACT_TOOL']=str(artifact)
        if not node:raise RuntimeError('Artifact authoring needs Node; set VALUATION_NODE')
        proc=subprocess.run([node,str(Path(__file__).with_name('workbook.mjs')),str(out/'workbook-spec.json'),str(out)],
                            env=env,text=True,capture_output=True,timeout=240)
        (out/'workbook-engine.log').write_text(proc.stdout+'\n'+proc.stderr)
        if proc.returncode:raise RuntimeError('Artifact authoring failed; see workbook-engine.log')
        return {'authoring_engine':'artifact-tool','engine_checks':'artifact-engine-checks.json','version':json.loads((artifact.parent.parent/'package.json').read_text()).get('version') if artifact and (artifact.parent.parent/'package.json').is_file() else 'unknown'}
    openpyxl_export(spec,out)
    return {'authoring_engine':'openpyxl','engine_checks':None,'note':'Formula authoring only; recalculation checked separately.'}


def audit(path,graph,require_cached=False):
    from openpyxl import load_workbook
    formulas=load_workbook(path,data_only=False)
    values=load_workbook(path,data_only=True)
    errors=[];missing=[];mismatch=[];formula_count=0;literalized=[]
    for ws in values:
        for row in ws:
            for cell in row:
                if cell.data_type=='e':errors.append(f'{ws.title}!{cell.coordinate}: {cell.value}')
    for node in graph.nodes.values():
        if node['metadata']:continue
        formula_count+=1
        if formulas[node['sheet']][node['cell']].data_type!='f':literalized.append(node['path'])
        value=values[node['sheet']][node['cell']].value
        if value is None:missing.append(node['path'])
        elif not isinstance(value,(int,float)) or not math.isclose(value,node['value'],rel_tol=1e-8,abs_tol=1e-7):
            mismatch.append(dict(path=node['path'],expected=node['value'],actual=value))
    formulas.close();values.close()
    return dict(formula_cells=formula_count,formula_errors=errors,literalized=literalized,missing_cached=missing,
                differences=mismatch,cached_values_required=require_cached,calculated_values_verified=not(errors or missing or mismatch or literalized),passed=not(errors or literalized or mismatch or (require_cached and missing)),
                scope='All calculation cells and all stored Excel errors; cached values do not establish source correctness.')


def recalculate(out,graph,mode='auto'):
    executable=os.environ.get('VALUATION_SOFFICE') or shutil.which('soffice') or shutil.which('libreoffice') or bundled('bin/override/soffice')
    if mode=='skip':return dict(status='not_requested',engine=None,financial_review='not_performed')
    if not executable:
        return dict(status='unavailable',engine='LibreOffice',reason='No executable found',financial_review='not_performed')
    try:
        version=subprocess.run([executable,'--version'],capture_output=True,text=True,timeout=20)
        if version.returncode:raise RuntimeError(version.stderr or version.stdout)
        with tempfile.TemporaryDirectory(prefix='.recalc-',dir=out) as temp:
            root=Path(temp);source=root/'input';dest=root/'output';source.mkdir();dest.mkdir()
            shutil.copy2(out/'valuation.xlsx',source/'valuation.xlsx')
            proc=subprocess.run([executable,'-env:UserInstallation='+(root/'profile').as_uri(),'--headless',
                                 '--convert-to','xlsx','--outdir',str(dest),str(source/'valuation.xlsx')],
                                capture_output=True,text=True,timeout=90)
            (out/'libreoffice.log').write_text(proc.stdout+'\n'+proc.stderr)
            candidate=dest/'valuation.xlsx'
            if proc.returncode or not candidate.exists():raise RuntimeError('Conversion did not produce a workbook; see libreoffice.log')
            checks=audit(candidate,graph,require_cached=True)
            if checks['passed']:shutil.copy2(candidate,out/'valuation.xlsx')
            else:shutil.copy2(candidate,out/'failed-recalculation.xlsx')
            return dict(status='passed' if checks['passed'] else 'failed',engine='LibreOffice',version=version.stdout.strip(),
                        recalculated=True,saved=True,reopened=True,checks=checks,financial_review='not_performed')
    except (OSError,subprocess.TimeoutExpired,RuntimeError) as exc:
        return dict(status='failed',engine='LibreOffice',reason=str(exc),financial_review='not_performed')
